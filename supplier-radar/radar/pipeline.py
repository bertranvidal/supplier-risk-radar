"""Procesa un DataFrame de noticias (en paralelo, con cache) y escribe el Excel de resultados."""
import hashlib
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

from . import llm
from .analyze import SYSTEM, Context, analyze
from .config import CACHE, TRAIN, VALID
from .data import load_news

REQUIRED = ["News_ID", "Is_Risk", "Urgent", "Affected_Supplier_IDs"]


def load_pool(include_valid: bool) -> pd.DataFrame:
    """Noticias etiquetadas que pueden usarse como ejemplos. Nunca se usan test ni live."""
    pool = load_news(TRAIN)
    if include_valid:
        pool = pd.concat([pool, load_news(VALID)], ignore_index=True)
    return pool


def _cache_key(mode: str, prompt: str) -> str:
    return hashlib.sha1((mode + "|" + llm.model_name() + "|" + SYSTEM + "|" + prompt).encode()).hexdigest()


def _one(row, ctx, mode, use_cache):
    from .analyze import build_prompt

    use_cache = use_cache and mode == "llm"  # el modo rules es instantaneo, no se cachea
    prompt = build_prompt(row, ctx) if mode == "llm" else ""
    path = CACHE / (_cache_key(mode, prompt) + ".json")
    if use_cache and path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    res, _ = analyze(row, ctx, mode)
    if use_cache:
        CACHE.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
    return res


def process(news: pd.DataFrame, pool: pd.DataFrame, mode: str = "llm", workers: int | None = None,
            use_cache: bool = True, verbose: bool = True) -> pd.DataFrame:
    ctx = Context.build(pool)
    workers = workers or int(os.getenv("RADAR_WORKERS", "4"))
    t0 = time.time()
    rows = [r for _, r in news.iterrows()]

    def work(r):
        try:
            res = _one(r, ctx, mode, use_cache)
        except Exception as e:  # noqa: BLE001 - nunca romper la demo por una noticia
            res = {"Is_Risk": "N", "Urgent": "N", "Affected_Supplier_IDs": "", "Event_Type": "ERROR",
                   "Supply_Impact": "", "Duration": "", "Confidence": "low", "Reasoning": "ERROR: " + str(e)[:200]}
        if verbose:
            print(f"  {r.News_ID}: Is_Risk={res['Is_Risk']} Urgent={res['Urgent']} {res['Affected_Supplier_IDs']}", flush=True)
        return {"News_ID": r.News_ID, **res}

    if mode == "llm" and workers > 1:
        with ThreadPoolExecutor(workers) as ex:
            results = list(ex.map(work, rows))
    else:
        results = [work(r) for r in rows]
    out = pd.DataFrame(results)
    out.insert(1, "Headline", news["Headline"].values)
    out.attrs["seconds"] = time.time() - t0
    if verbose:
        print(f"Procesadas {len(out)} noticias en {out.attrs['seconds']:.1f} s")
    return out


def analyze_news(news: pd.DataFrame, pool: pd.DataFrame, mode: str = "models", alerts: bool = False,
                 alert_lang: str = "en", alert_llm: bool = True, **kw) -> pd.DataFrame:
    """process() + capa de impacto de negocio (+ alertas redactadas con IA generativa si alerts=True)."""
    from .data import load_materials, load_suppliers
    from .impact import enrich

    res = process(news, pool, mode, **kw)
    out = enrich(res, news, load_suppliers(), load_materials())
    out.attrs["seconds"] = res.attrs["seconds"]
    if alerts:
        from .alerts import add_alerts

        t0 = time.time()
        out = add_alerts(out, news, load_suppliers(), lang=alert_lang, use_llm=alert_llm)
        out.attrs["seconds"] = res.attrs["seconds"]
        out.attrs["alert_seconds"] = time.time() - t0
    return out


def save_excel(df: pd.DataFrame, path) -> None:
    """Hoja 'Results' con SOLO las 4 columnas pedidas y hoja 'Details' con todo."""
    path = str(path) if isinstance(path, (str, Path)) else path  # tambien vale un BytesIO
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df[REQUIRED].to_excel(xw, sheet_name="Results", index=False)
        df.drop(columns=[]).to_excel(xw, sheet_name="Details", index=False)
        red = PatternFill("solid", fgColor="F8CBAD")
        amber = PatternFill("solid", fgColor="FFE699")
        for name in ("Results", "Details"):
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
            hdr = [c.value for c in ws[1]]
            ir, iu = hdr.index("Is_Risk"), hdr.index("Urgent")
            for row in ws.iter_rows(min_row=2):
                fill = red if row[iu].value == "Y" else (amber if row[ir].value == "Y" else None)
                if fill:
                    for c in row:
                        c.fill = fill
            for col in ws.columns:
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
